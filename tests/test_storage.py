import csv
import os
import tempfile
import unittest

from pingaroo.storage import CSV_COLUMNS, StateFile, Storage
from tests.fakes import test_config


class StorageTests(unittest.TestCase):
    def test_csv_header_and_rotation(self):
        with tempfile.TemporaryDirectory() as d:
            cfg = test_config()
            cfg["logging"]["directory"] = d
            cfg["logging"]["csv_max_bytes"] = 600
            cfg["logging"]["retention_files"] = 2
            s = Storage(cfg)
            for i in range(12):
                self.assertTrue(s.write_result_row({"cycle_id": f"c{i}", "status": "SUCCESS", "error": "x" * 60}))
            files = sorted(os.listdir(d))
            rotated = [f for f in files if f.startswith("results.") and f != "results.csv"]
            self.assertLessEqual(len(rotated), 2)
            with open(s.csv_path, newline="") as fh:
                rows = list(csv.reader(fh))
            self.assertEqual(rows[0], CSV_COLUMNS)
            for f in rotated:
                with open(os.path.join(d, f), newline="") as fh:
                    self.assertEqual(next(csv.reader(fh)), CSV_COLUMNS)

    def test_write_error_reported(self):
        with tempfile.TemporaryDirectory() as d:
            cfg = test_config()
            cfg["logging"]["directory"] = os.path.join(d, "nope")
            s = Storage(cfg)
            with open(cfg["logging"]["directory"], "w") as fh:
                fh.write("blocker")
            self.assertFalse(s.write_event("x"))
            self.assertIn("cannot create", s.status.last_error)

    def test_recovery_clears_error(self):
        with tempfile.TemporaryDirectory() as d:
            cfg = test_config()
            cfg["logging"]["directory"] = d
            s = Storage(cfg)
            s.status.last_error = "old"
            self.assertTrue(s.write_event("ok"))
            self.assertEqual(s.status.last_error, "")

    def test_state_and_control_file(self):
        with tempfile.TemporaryDirectory() as d:
            sf = StateFile(d)
            sf.write({"a": 1})
            self.assertEqual(sf.read()["a"], 1)
            self.assertEqual(sf.pop_control(), "")
            sf.push_control("start")
            self.assertEqual(sf.pop_control(), "start")
            self.assertEqual(sf.pop_control(), "")


if __name__ == "__main__":
    unittest.main()


class ExportTests(unittest.TestCase):
    def test_xlsx_export_and_summary(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        from pingaroo import export
        with tempfile.TemporaryDirectory() as d:
            cfg = test_config()
            cfg["logging"]["directory"] = d
            s = Storage(cfg)
            s.write_result_row({"cycle_id": "a-1", "label": "Config A", "status": "SUCCESS", "download_mbps": 50.0,
                                "upload_mbps": 10.0, "latency_ms": 30.0, "start_time": "2026-10-29T10:00:00-05:00"})
            s.write_result_row({"cycle_id": "a-2", "label": "Config A", "status": "SPEEDTEST_FAILED",
                                "start_time": "2026-10-29T10:05:00-05:00", "error": "x"})
            s.write_result_row({"cycle_id": "b-1", "label": "Config B", "status": "SUCCESS", "download_mbps": 70.0,
                                "upload_mbps": 20.0, "latency_ms": 40.0, "start_time": "2026-10-29T10:10:00-05:00"})
            rows = export.load_rows(d)
            self.assertEqual(len(rows), 3)
            summary = export.summarize(rows)
            a = next(x for x in summary if x["label"] == "Config A")
            self.assertEqual((a["attempts"], a["successes"], a["success_pct"]), (2, 1, 50.0))
            self.assertEqual(a["down_avg"], 50.0)
            self.assertIn("SPEEDTEST_FAILED=1", a["failures"])
            out = os.path.join(d, "x.xlsx")
            export.write_xlsx(rows, out)
            wb = openpyxl.load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Results", "Summary by Label", "About"])
            self.assertEqual(wb["Results"].max_row, 4)
            self.assertEqual(wb["Summary by Label"].max_row, 3)
