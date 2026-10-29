"""Excel export of results with a per-label summary sheet (openpyxl)."""

from __future__ import annotations

import csv
import glob
import os
import statistics
import time

from .storage import CSV_COLUMNS

NUMERIC = {"duration_s", "signal_dbm", "channel", "link_rate_mbps", "dns_ms", "https_status", "https_ms",
           "gateway_ping_ms", "gateway_ping_loss_pct", "ping1_avg_ms", "ping1_loss_pct", "ping2_avg_ms",
           "ping2_loss_pct", "download_mbps", "upload_mbps", "latency_ms", "jitter_ms", "packet_loss_pct",
           "download_bytes", "upload_bytes", "server_id", "cell_rsrp", "cell_rsrq", "cell_sinr", "cell_rssi",
           "cell_earfcn", "cell_pci", "cell_tac", "cell_cell_id", "cell_5g_rsrp", "cell_5g_rsrq", "cell_5g_sinr",
           "cell_5g_pci", "cell_bandwidth_mhz"}


def _num(v):
    if v in ("", None):
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in str(v) else f
    except ValueError:
        return v


def load_rows(directory: str) -> list[dict]:
    files = sorted(glob.glob(os.path.join(directory, "results.*.csv"))) + [os.path.join(directory, "results.csv")]
    rows: list[dict] = []
    for path in files:
        if not os.path.exists(path):
            continue
        with open(path, newline="", encoding="utf-8") as fh:
            rows.extend(csv.DictReader(fh))
    rows.sort(key=lambda r: r.get("start_time", ""))
    return rows


def summarize(rows: list[dict]) -> list[dict]:
    by_label: dict[str, list[dict]] = {}
    for r in rows:
        by_label.setdefault(r.get("label", ""), []).append(r)
    out = []
    for label, group in sorted(by_label.items()):
        ok = [r for r in group if r.get("status") == "SUCCESS"]

        def stats(col):
            vals = [float(r[col]) for r in ok if r.get(col) not in ("", None)]
            if not vals:
                return (None, None, None, None)
            return (round(statistics.mean(vals), 2), round(statistics.median(vals), 2), round(min(vals), 2), round(max(vals), 2))

        d = stats("download_mbps")
        u = stats("upload_mbps")
        l = stats("latency_ms")
        s = stats("signal_dbm")
        rsrp = stats("cell_rsrp")
        statuses = {}
        for r in group:
            statuses[r.get("status", "")] = statuses.get(r.get("status", ""), 0) + 1
        out.append({
            "label": label, "attempts": len(group), "successes": len(ok),
            "success_pct": round(100.0 * len(ok) / len(group), 1) if group else None,
            "first": group[0].get("start_time", ""), "last": group[-1].get("end_time", ""),
            "down_avg": d[0], "down_median": d[1], "down_min": d[2], "down_max": d[3],
            "up_avg": u[0], "up_median": u[1], "up_min": u[2], "up_max": u[3],
            "latency_avg": l[0], "latency_median": l[1], "latency_min": l[2], "latency_max": l[3],
            "wifi_dbm_avg": s[0], "rsrp_avg": rsrp[0],
            "failures": "; ".join(f"{k}={v}" for k, v in sorted(statuses.items()) if k != "SUCCESS"),
        })
    return out


def write_xlsx(rows: list[dict], path: str, title: str = "Pingaroo results") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    ws.append(CSV_COLUMNS)
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([(_num(r.get(c)) if c in NUMERIC else r.get(c, "")) for c in CSV_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"start_time": 26, "end_time": 26, "error": 48, "cycle_id": 26, "session_id": 22, "label": 16,
              "status": 20, "result_url": 50, "server_name": 20, "server_location": 26, "https_url": 40}
    for idx, col in enumerate(CSV_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, max(10, min(18, len(col) + 2)))
    status_fill = {"SUCCESS": "C6EFCE", "INTERRUPTED": "FFEB9C"}
    status_col = CSV_COLUMNS.index("status") + 1
    for row in ws.iter_rows(min_row=2, min_col=status_col, max_col=status_col):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=status_fill.get(cell.value, "FFC7CE"))

    ws2 = wb.create_sheet("Summary by Label")
    cols = ["label", "attempts", "successes", "success_pct", "first", "last", "down_avg", "down_median", "down_min",
            "down_max", "up_avg", "up_median", "up_min", "up_max", "latency_avg", "latency_median", "latency_min",
            "latency_max", "wifi_dbm_avg", "rsrp_avg", "failures"]
    ws2.append(cols)
    for cell in ws2[1]:
        cell.font = head_font
        cell.fill = head_fill
    for srow in summarize(rows):
        ws2.append([srow.get(c) for c in cols])
    ws2.freeze_panes = "A2"
    for idx, col in enumerate(cols, start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = 26 if col in ("first", "last") else 40 if col == "failures" else 14
    ws3 = wb.create_sheet("About")
    ws3.append(["Pingaroo export", time.strftime("%Y-%m-%d %H:%M:%S %z")])
    ws3.append(["Rows", len(rows)])
    ws3.append(["Units", "Mbps = megabits per second; latency/jitter ms; RSRP/RSRQ dBm/dB; signal_dbm is the Pi's Wi-Fi link, not cellular"])
    ws3.append(["Note", "Measured throughput includes the Pi-to-router Wi-Fi link"])
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
